"""Koikatsu / Koikatsu Sunshine card and zipmod parsing.

Card files are PNGs with an Illusion payload appended after IEND. The payload
contains MessagePack blocks including KKEx, which lists the zipmod GUIDs the
card needs. zipmods are zip files with a manifest.xml that declares a GUID.
"""

from __future__ import annotations

import csv
import io
import json
import os
import re
import struct
import subprocess
import tempfile
import zipfile
from collections import defaultdict
from pathlib import Path
from xml.etree import ElementTree as ET

try:
    import msgpack
except ImportError:
    msgpack = None

PNG_END = b"IEND\xaeB`\x82"
CHARACTER_TYPES = {
    "【KoiKatuChara】",
    "【KoiKatuCharaS】",
    "【KoiKatuCharaSP】",
    "【KoiKatuCharaSun】",
}
CLOTHES_TYPE = "【KoiKatuClothes】"
SCENE_MARKER = b"\x04KKEx"


def _require_msgpack():
    if msgpack is None:
        raise RuntimeError("需要 msgpack: python -m pip install msgpack")


def _unpack(blob):
    _require_msgpack()
    if not blob:
        return None
    try:
        return msgpack.unpackb(blob, raw=False, strict_map_key=False)
    except Exception:
        return None


def _read_7bit(data: bytes, pos: int):
    result = 0
    shift = 0
    while True:
        if pos >= len(data) or shift >= 35:
            raise ValueError("bad 7-bit integer")
        byte = data[pos]
        pos += 1
        result |= (byte & 0x7F) << shift
        if not byte & 0x80:
            return result, pos
        shift += 7


def _read_string(data: bytes, pos: int):
    length, pos = _read_7bit(data, pos)
    if pos + length > len(data):
        raise ValueError("truncated string")
    return data[pos:pos + length].decode("utf-8", "ignore"), pos + length


def _looks_like_version(value: str):
    if not value:
        return False
    parts = value.split(".")
    return len(parts) >= 2 and all(part.isdigit() for part in parts)


def _bytes_key(value):
    if isinstance(value, bytes):
        return value.decode("utf-8", "ignore")
    return str(value)


def parse_plugin_map(blob: bytes) -> list[dict]:
    packed = _unpack(blob)
    if not isinstance(packed, dict):
        return []
    found: list[dict] = []
    seen = set()
    for plugin_name, fields in packed.items():
        if not isinstance(fields, list) or len(fields) < 2:
            continue
        data = fields[1]
        if not isinstance(data, dict):
            continue
        info = data.get("info")
        if not isinstance(info, list):
            continue
        for entry in info:
            if not isinstance(entry, (bytes, bytearray)):
                continue
            resolve = _unpack(bytes(entry))
            if not isinstance(resolve, dict):
                continue
            guid = resolve.get("ModID") or resolve.get("GUID") or resolve.get("guid")
            if not guid:
                continue
            guid = _bytes_key(guid)
            prop = _bytes_key(resolve.get("Property", ""))
            key = (guid, prop)
            if key in seen:
                continue
            seen.add(key)
            found.append({
                "guid": guid,
                "property": prop,
                "slot": resolve.get("Slot", 0),
                "local_slot": resolve.get("LocalSlot", 0),
                "plugin": _bytes_key(plugin_name),
            })
    return found


def _read_blocks(payload: bytes, pos: int):
    if pos + 4 > len(payload):
        raise ValueError("truncated block index length")
    index_len = struct.unpack_from("<I", payload, pos)[0]
    pos += 4
    if pos + index_len > len(payload):
        raise ValueError("truncated block index")
    index = _unpack(payload[pos:pos + index_len])
    pos += index_len
    if not isinstance(index, dict):
        raise ValueError("bad block index")
    entries = index.get("lstInfo")
    if not isinstance(entries, list):
        raise ValueError("missing lstInfo")
    if pos + 8 > len(payload):
        raise ValueError("truncated block data length")
    raw_len = struct.unpack_from("<Q", payload, pos)[0]
    pos += 8
    if pos + raw_len > len(payload):
        raise ValueError("truncated block data")
    raw = payload[pos:pos + raw_len]
    pos += raw_len
    blocks = {}
    for entry in entries:
        if not isinstance(entry, dict):
            continue
        name = _bytes_key(entry.get("name", ""))
        if not name:
            continue
        start = int(entry.get("pos", 0) or 0)
        size = int(entry.get("size", 0) or 0)
        if start < 0 or size < 0 or start + size > len(raw):
            continue
        blocks[name] = {
            "version": _bytes_key(entry.get("version", "")),
            "data": raw[start:start + size],
        }
    return blocks, pos


def _character_info(blocks: dict) -> dict:
    parameter = blocks.get("Parameter")
    if not parameter:
        return {}
    data = _unpack(parameter["data"])
    if not isinstance(data, dict):
        return {}
    result = {}
    for key in ("lastname", "firstname", "nickname", "sex", "LastName", "FirstName", "Nickname", "Sex"):
        if key in data and data[key] not in (None, ""):
            result[key.lower()] = data[key]
    return result


def _parse_character(payload: bytes, pos: int):
    version, pos = _read_string(payload, pos)
    if pos + 4 > len(payload):
        raise ValueError("truncated face length")
    face_len = struct.unpack_from("<I", payload, pos)[0]
    pos += 4 + face_len
    blocks, _ = _read_blocks(payload, pos)
    required = parse_plugin_map(blocks.get("KKEx", {}).get("data", b""))
    return {
        "type": "character",
        "version": version,
        "required": required,
        "character": _character_info(blocks),
        "blocks": sorted(blocks),
    }


def _parse_clothes(payload: bytes, pos: int):
    version, pos = _read_string(payload, pos)
    coordinate, pos = _read_string(payload, pos)
    if pos + 4 > len(payload):
        raise ValueError("truncated clothes header")
    length = struct.unpack_from("<i", payload, pos)[0]
    pos += 4
    if length > 0:
        pos += length
    marker, pos = _read_string(payload, pos)
    version2, pos = _read_string(payload, pos)
    if pos + 4 > len(payload):
        raise ValueError("truncated clothes KKEx length")
    kkex_len = struct.unpack_from("<i", payload, pos)[0]
    pos += 4
    blob = payload[pos:pos + max(0, kkex_len)]
    required = parse_plugin_map(blob) if marker == "KKEx" else []
    return {
        "type": "clothes",
        "version": version,
        "coordinate": coordinate,
        "required": required,
        "blocks": [],
    }


def _dedupe_required(items: list[dict]) -> list[dict]:
    seen = set()
    result = []
    for item in items:
        key = (item.get("guid"), item.get("property", ""))
        if key in seen:
            continue
        seen.add(key)
        result.append(item)
    return result


def _parse_scene(payload: bytes):
    version, _ = _read_string(payload, 0)
    required: list[dict] = []
    extended_version = None
    search_end = len(payload)
    while search_end > 0:
        index = payload.rfind(SCENE_MARKER, 0, search_end)
        if index < 0:
            break
        header_end = index + len(SCENE_MARKER) + 8
        if header_end <= len(payload):
            ext_version, length = struct.unpack_from("<ii", payload, index + len(SCENE_MARKER))
            data_end = header_end + length
            if ext_version >= 0 and length > 0 and data_end == len(payload):
                required = parse_plugin_map(payload[header_end:data_end])
                extended_version = ext_version
                break
        search_end = index

    embedded = 0
    for card_type in CHARACTER_TYPES:
        encoded = card_type.encode("utf-8")
        if len(encoded) >= 128:
            continue
        marker = bytes([len(encoded)]) + encoded
        start = 0
        while True:
            index = payload.find(marker, start)
            if index < 0:
                break
            start = index + len(marker)
            try:
                info = _parse_character(payload, start)
                required.extend(info["required"])
                embedded += 1
            except Exception:
                continue
    return {
        "type": "scene",
        "version": version,
        "required": _dedupe_required(required),
        "embedded_characters": embedded,
        "extended_version": extended_version,
        "blocks": [],
    }


def parse_card(path: Path) -> dict | None:
    try:
        data = Path(path).read_bytes()
        end = data.find(PNG_END)
        if end < 0:
            return None
        payload = data[end + len(PNG_END):]
        if not payload:
            return None
        first = payload[0]
        pos = 0
        if first == 0x07:
            card_type = "Scene"
        elif first == 0x64:
            pos = 4
            card_type, pos = _read_string(payload, pos)
        else:
            candidate, _ = _read_string(payload, 0)
            if _looks_like_version(candidate):
                card_type = "Scene"
            elif candidate in CHARACTER_TYPES or candidate == CLOTHES_TYPE:
                card_type = candidate
                pos = len(payload) - len(payload)  # type starts at 0; parse below
                _, pos = _read_string(payload, 0)
            else:
                return None

        if card_type == "Scene":
            info = _parse_scene(payload)
        elif card_type == CLOTHES_TYPE:
            info = _parse_clothes(payload, pos)
        elif card_type in CHARACTER_TYPES:
            info = _parse_character(payload, pos)
        else:
            return None
        info["card_type"] = card_type
        info["source"] = str(path)
        return info
    except Exception:
        return None


def _manifest_from_zip(archive: zipfile.ZipFile) -> dict | None:
    names = [name for name in archive.namelist() if name.lower().endswith("manifest.xml")]
    guids: list[str] = []
    result: dict = {}
    for name in names[:8]:
        try:
            root = ET.fromstring(archive.read(name))
        except Exception:
            continue
        fields = {}
        for element in root.iter():
            tag = element.tag.split("}")[-1].lower()
            if tag in {"guid", "name", "version", "author"} and element.text:
                fields.setdefault(tag, element.text.strip())
        guid = fields.get("guid")
        if guid and guid not in guids:
            guids.append(guid)
        for key in ("name", "author", "version"):
            if fields.get(key) and not result.get(key):
                result[key] = fields[key]
    if not guids:
        return None
    result["guids"] = guids
    result["type"] = "zipmod"
    return result


def parse_zipmod(path: Path) -> dict | None:
    path = Path(path)
    try:
        with zipfile.ZipFile(path) as archive:
            manifest = _manifest_from_zip(archive)
            if manifest:
                return manifest
            for name in archive.namelist():
                if not name.lower().endswith(".zipmod"):
                    continue
                try:
                    with zipfile.ZipFile(io.BytesIO(archive.read(name))) as nested:
                        manifest = _manifest_from_zip(nested)
                        if manifest:
                            return manifest
                except Exception:
                    continue
    except Exception:
        return None
    return None


def scan_folder(path: Path, card_limit: int = 12, mod_limit: int = 12) -> dict | None:
    path = Path(path)
    if not path.is_dir():
        return None
    cards: list[dict] = []
    mods: list[dict] = []
    try:
        for dirpath, dirnames, filenames in os.walk(path):
            depth = len(Path(dirpath).relative_to(path).parts)
            if depth > 3:
                dirnames[:] = []
                continue
            for filename in filenames:
                entry = Path(dirpath) / filename
                suffix = entry.suffix.lower()
                if suffix == ".png" and len(cards) < card_limit:
                    info = parse_card(entry)
                    if info:
                        cards.append(info)
                elif suffix == ".zipmod" and len(mods) < mod_limit:
                    info = parse_zipmod(entry)
                    if info:
                        mods.append(info)
                if len(cards) >= card_limit and len(mods) >= mod_limit:
                    break
            if len(cards) >= card_limit and len(mods) >= mod_limit:
                break
    except OSError:
        return None
    if not cards and not mods:
        return None
    required = []
    for card in cards:
        required.extend(card.get("required", []))
    guids = []
    for mod in mods:
        guids.extend(mod.get("guids", []))
    return {
        "type": "pack",
        "card_types": sorted({card["type"] for card in cards}),
        "cards": len(cards),
        "mods": len(mods),
        "has_mods": bool(mods),
        "required": _dedupe_required(required),
        "guids": sorted(set(guids)),
    }


def _load_sidecar(path: Path) -> dict | None:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return None


def link_library(root: Path, apply: bool = False, shortcuts: bool = False) -> dict:
    root = Path(root)
    base = root / "11_Koikatsu"
    cards: list[tuple[Path, dict]] = []
    mods: list[tuple[Path, dict]] = []
    for sidecar in base.rglob("_asset.json"):
        data = _load_sidecar(sidecar)
        if not data:
            continue
        koi = data.get("koi") or {}
        kind = koi.get("type")
        if kind in {"character", "clothes", "scene", "pack"}:
            cards.append((sidecar.parent, data))
        elif kind == "zipmod":
            mods.append((sidecar.parent, data))

    guid_index: dict[str, list[tuple[str, Path]]] = defaultdict(list)
    for folder, data in mods:
        koi = data.get("koi") or {}
        for guid in koi.get("guids", []):
            guid_index[guid].append((data.get("title") or folder.name, folder))

    rows = []
    shortcut_map: dict[tuple[Path, Path], tuple[str, str]] = {}
    missing_map: dict[Path, list[str]] = defaultdict(list)
    for folder, data in cards:
        koi = data.get("koi") or {}
        required = koi.get("required", [])
        lines = [
            f"卡片: {data.get('title') or folder.name}",
            f"类型: {koi.get('type')}  {koi.get('card_type', '')}".strip(),
            f"路径: {folder}",
            f"需要 MOD: {len(required)}",
            "",
        ]
        installed = 0
        missing = 0
        for item in required:
            guid = item.get("guid", "")
            hits = guid_index.get(guid, [])
            if hits:
                installed += 1
                for name, mod_folder in hits:
                    lines.append(f"[已安装] {guid} | {name} | {mod_folder}")
                    rows.append({
                        "卡片": data.get("title") or folder.name,
                        "类型": koi.get("type", ""),
                        "GUID": guid,
                        "状态": "已安装",
                        "MOD名称": name,
                        "MOD路径": str(mod_folder),
                        "卡片路径": str(folder),
                    })
                    shortcut_map.setdefault((folder, mod_folder), (name, guid))
            else:
                missing += 1
                lines.append(f"[缺失]   {guid}")
                missing_map[folder].append(guid)
                rows.append({
                    "卡片": data.get("title") or folder.name,
                    "类型": koi.get("type", ""),
                    "GUID": guid,
                    "状态": "缺失",
                    "MOD名称": "",
                    "MOD路径": "",
                    "卡片路径": str(folder),
                })
        if not required:
            lines.append("无 MOD 依赖。")
        lines.extend(["", f"汇总: 已安装 {installed} / 缺失 {missing}"])
        if apply:
            (folder / "_mod_links.txt").write_text("\n".join(lines), encoding="utf-8")

    shortcut_count = 0
    if apply and shortcuts and shortcut_map:
        specs = []
        for (card_folder, mod_folder), (name, guid) in shortcut_map.items():
            mods_dir = card_folder / "_mods"
            mods_dir.mkdir(exist_ok=True)
            safe = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name).strip(" .") or mod_folder.name
            target = mods_dir / f"{safe}.lnk"
            specs.append((target, mod_folder, guid))
        if specs:
            script = ["$ws = New-Object -ComObject WScript.Shell"]
            for target, mod_folder, guid in specs:
                def ps(value):
                    return str(value).replace("'", "''")
                script.append(f"$s = $ws.CreateShortcut('{ps(target)}')")
                script.append(f"$s.TargetPath = '{ps(mod_folder)}'")
                script.append(f"$s.WorkingDirectory = '{ps(mod_folder)}'")
                script.append(f"$s.Description = '{ps(guid)}'")
                script.append("$s.Save()")
            script_path = Path(tempfile.gettempdir()) / "koi_shortcuts.ps1"
            script_path.write_text("\n".join(script), encoding="utf-8-sig")
            try:
                subprocess.run(
                    ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(script_path)],
                    capture_output=True, timeout=120,
                )
                shortcut_count = len(specs)
            except Exception:
                shortcut_count = 0
    for folder, guids in missing_map.items():
        if apply:
            (folder / "_missing_mods.txt").write_text("\n".join(sorted(set(guids))), encoding="utf-8")

    summary = {
        "cards": len(cards),
        "mods": len(mods),
        "rows": len(rows),
        "installed": sum(1 for row in rows if row["状态"] == "已安装"),
        "missing": sum(1 for row in rows if row["状态"] == "缺失"),
        "shortcuts": shortcut_count,
    }
    if apply:
        index_dir = base / "_Index"
        index_dir.mkdir(parents=True, exist_ok=True)
        csv_path = index_dir / "卡片MOD依赖.csv"
        with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=["卡片", "类型", "GUID", "状态", "MOD名称", "MOD路径", "卡片路径"])
            writer.writeheader()
            writer.writerows(rows)
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            wb = Workbook()
            ws = wb.active
            ws.title = "卡片MOD依赖"
            headers = ["卡片", "类型", "GUID", "状态", "MOD名称", "MOD路径", "卡片路径"]
            for column, text in enumerate(headers, 1):
                cell = ws.cell(row=1, column=column, value=text)
                cell.fill = PatternFill("solid", fgColor="0F766E")
                cell.font = Font(color="FFFFFF", bold=True)
            link_font = Font(color="0563C1", underline="single")
            for index, row in enumerate(rows, 2):
                for column, key in enumerate(headers, 1):
                    cell = ws.cell(row=index, column=column, value=row[key])
                    if key in {"MOD路径", "卡片路径"} and row[key]:
                        cell.hyperlink = row[key]
                        cell.font = link_font
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = f"A1:G{max(1, len(rows) + 1)}"
            for column, width in zip("ABCDEFG", (28, 10, 40, 10, 24, 60, 60)):
                ws.column_dimensions[column].width = width
            wb.save(index_dir / "卡片MOD依赖.xlsx")
        except Exception:
            pass
    return summary


if __name__ == "__main__":
    import argparse
    import os
    import sys

    parser = argparse.ArgumentParser()
    parser.add_argument("path")
    parser.add_argument("--type", choices=["auto", "card", "zipmod"], default="auto")
    args = parser.parse_args()
    p = Path(args.path)
    if args.type in {"auto", "card"} and p.suffix.lower() == ".png":
        print(json.dumps(parse_card(p), ensure_ascii=False, indent=2))
    elif args.type in {"auto", "zipmod"} and p.suffix.lower() == ".zipmod":
        print(json.dumps(parse_zipmod(p), ensure_ascii=False, indent=2))
    else:
        print(json.dumps(scan_folder(p), ensure_ascii=False, indent=2))
